"""
Data Export Service for Retail SKU-Week Demand Forecasting.
Generates compliant, verified CSV and XLSX export workbooks from persisted forecast records.
Enforces atomic writes, cryptographic SHA-256 tracking, and path containment security.
"""
import csv
import hashlib
from pathlib import Path
import tempfile
from typing import Any, Dict, List, Optional
import uuid

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from src.core.config import settings
from src.core.logging_config import logger
from src.api.errors import (
    ExportFileMissingOrCorruptException,
    InvalidExportFormatException,
    RunNotCompletedException,
)
from src.api.schemas.forecast import (
    ExportFormat,
    ExportStatus,
    ExportArtifactResponse,
    ForecastRunStatus,
)
from src.services.persistence_service import PersistenceService, persistence_service


CSV_COLUMNS = [
    "forecast_run_id",
    "model_id",
    "forecast_origin",
    "horizon_step",
    "target_week",
    "sku_id",
    "prediction",
    "prediction_raw",
    "prediction_status",
    "clipped_at_zero",
]


def sanitize_csv_cell(value: Any) -> str:
    """
    Neutralizes spreadsheet formula injection (CSV/DDE injection) attacks.
    Prefixes cells starting with '=', '+', '-', '@', tab or CR with a single quote.
    """
    s = str(value) if value is not None else ""
    if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{s}"
    return s


class ExportService:
    """Service generating downloadable exports (CSV/XLSX) from persisted forecast runs."""

    def __init__(self, persistence: Optional[PersistenceService] = None, exports_dir: Optional[Path] = None):
        self.persistence = persistence or persistence_service
        self.exports_dir = exports_dir or settings.EXPORTS_DIR
        self.exports_dir.mkdir(parents=True, exist_ok=True)

    def generate_export(
        self,
        forecast_run_id: str,
        format_type: ExportFormat,
        user_id: Optional[str] = None,
    ) -> ExportArtifactResponse:
        """
        Builds export file from persisted records.
        Rejects non-SUCCEEDED runs. Atomically writes file and records metadata.
        """
        run = self.persistence.get_run(forecast_run_id)
        if run.status != ForecastRunStatus.SUCCEEDED:
            raise RunNotCompletedException(forecast_run_id, run.status.value)

        records = self.persistence.get_all_run_records_for_export(forecast_run_id)
        export_id = f"exp_{uuid.uuid4().hex[:12]}"

        if format_type == ExportFormat.CSV:
            ext = "csv"
            managed_file_key = f"{forecast_run_id}_{export_id}.{ext}"
            target_path = self.exports_dir / managed_file_key
            row_count, file_size, sha256_hash = self._write_csv(run, records, target_path)
        elif format_type == ExportFormat.XLSX:
            ext = "xlsx"
            managed_file_key = f"{forecast_run_id}_{export_id}.{ext}"
            target_path = self.exports_dir / managed_file_key
            row_count, file_size, sha256_hash = self._write_xlsx(run, records, target_path)
        else:
            raise InvalidExportFormatException(str(format_type), [e.value for e in ExportFormat])

        # Record export in database
        artifact = self.persistence.record_export(
            export_id=export_id,
            forecast_run_id=forecast_run_id,
            format_type=format_type,
            managed_file_key=managed_file_key,
            row_count=row_count,
            file_size_bytes=file_size,
            sha256_hash=sha256_hash,
            status=ExportStatus.COMPLETED,
            user_id=user_id,
        )
        logger.info(
            f"Generated {format_type.value} export '{export_id}' for run '{forecast_run_id}' "
            f"({row_count} rows, {file_size} bytes, sha256={sha256_hash[:12]})."
        )
        return artifact

    def _write_csv(
        self,
        run: Any,
        records: List[Dict[str, Any]],
        target_path: Path,
    ) -> tuple[int, int, str]:
        """Atomically writes CSV using standard RFC 4180 format with UTF-8 encoding."""
        with tempfile.NamedTemporaryFile("w", delete=False, newline="", encoding="utf-8") as tmp:
            tmp_path = Path(tmp.name)
            writer = csv.writer(tmp, quoting=csv.QUOTE_MINIMAL)
            writer.writerow(CSV_COLUMNS)

            for r in records:
                writer.writerow([
                    sanitize_csv_cell(run.forecast_run_id),
                    sanitize_csv_cell(run.model_id),
                    sanitize_csv_cell(run.forecast_origin),
                    r["horizon_step"],
                    sanitize_csv_cell(r["target_week"]),
                    sanitize_csv_cell(r["sku_id"]),  # Preserve string identity + formula injection defense
                    f"{float(r['prediction']):.4f}",
                    f"{float(r['prediction_raw']):.4f}" if r["prediction_raw"] is not None else "",
                    sanitize_csv_cell(r["prediction_status"]),
                    "TRUE" if bool(r["clipped_at_zero"]) else "FALSE",
                ])

        # Compute hash and move atomically
        content_bytes = tmp_path.read_bytes()
        sha256_hash = hashlib.sha256(content_bytes).hexdigest()
        file_size = len(content_bytes)
        tmp_path.replace(target_path)

        return len(records), file_size, sha256_hash

    def _write_xlsx(
        self,
        run: Any,
        records: List[Dict[str, Any]],
        target_path: Path,
    ) -> tuple[int, int, str]:
        """Writes Excel workbook using openpyxl with explicit formatting and metadata tab."""
        wb = openpyxl.Workbook()

        # Tab 1: Forecast Records
        ws_records = wb.active
        ws_records.title = "Forecast Records"

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")

        ws_records.append(CSV_COLUMNS)
        for col_idx in range(1, len(CSV_COLUMNS) + 1):
            cell = ws_records.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        for row_idx, r in enumerate(records, start=2):
            ws_records.cell(row=row_idx, column=1, value=str(run.forecast_run_id))
            ws_records.cell(row=row_idx, column=2, value=str(run.model_id))
            ws_records.cell(row=row_idx, column=3, value=str(run.forecast_origin))
            ws_records.cell(row=row_idx, column=4, value=int(r["horizon_step"]))
            ws_records.cell(row=row_idx, column=5, value=str(r["target_week"]))

            # Explicit string type for SKU ID to prevent loss of leading characters in spreadsheet apps
            sku_cell = ws_records.cell(row=row_idx, column=6, value=str(r["sku_id"]))
            sku_cell.data_type = "s"

            pred_cell = ws_records.cell(row=row_idx, column=7, value=round(float(r["prediction"]), 4))
            pred_cell.number_format = "0.00"

            if r["prediction_raw"] is not None:
                raw_cell = ws_records.cell(row=row_idx, column=8, value=round(float(r["prediction_raw"]), 4))
                raw_cell.number_format = "0.00"
            else:
                ws_records.cell(row=row_idx, column=8, value="")

            ws_records.cell(row=row_idx, column=9, value=str(r["prediction_status"]))
            ws_records.cell(row=row_idx, column=10, value="TRUE" if bool(r["clipped_at_zero"]) else "FALSE")

        # Tab 2: Run Metadata
        ws_meta = wb.create_sheet(title="Run Metadata")
        ws_meta.append(["Metadata Field", "Value"])
        ws_meta.cell(row=1, column=1).font = header_font
        ws_meta.cell(row=1, column=1).fill = header_fill
        ws_meta.cell(row=1, column=2).font = header_font
        ws_meta.cell(row=1, column=2).fill = header_fill

        metadata_items = [
            ("Forecast Run ID", run.forecast_run_id),
            ("Execution Status", run.status.value),
            ("Model ID", run.model_id),
            ("Model Version", run.model_version),
            ("Forecast Origin Cutoff", run.forecast_origin),
            ("Horizon (Weeks)", str(run.horizon_weeks)),
            ("Requested SKU Count", str(run.requested_sku_count)),
            ("Prediction Records Count", str(run.prediction_count)),
            ("Created UTC", run.created_at_utc),
            ("Data Snapshot ID", run.data_snapshot_id or "N/A"),
            ("Training Cutoff", run.training_cutoff or "N/A"),
        ]
        for field_name, val in metadata_items:
            ws_meta.append([field_name, val])

        # Write to temporary file and move atomically
        with tempfile.NamedTemporaryFile("wb", delete=False, suffix=".xlsx") as tmp:
            tmp_path = Path(tmp.name)
            wb.save(tmp_path)

        content_bytes = tmp_path.read_bytes()
        sha256_hash = hashlib.sha256(content_bytes).hexdigest()
        file_size = len(content_bytes)
        tmp_path.replace(target_path)

        return len(records), file_size, sha256_hash

    def get_export_file_path(self, forecast_run_id: str, export_id: str) -> Path:
        """
        Retrieves physical file path after verifying path containment and existence.
        Prevents directory traversal or arbitrary file download vulnerabilities.
        """
        # Strict parameter validation against directory traversal payloads
        if ".." in forecast_run_id or "/" in forecast_run_id or "\\" in forecast_run_id:
            raise ExportFileMissingOrCorruptException(export_id, "Invalid characters in forecast_run_id.")
        if ".." in export_id or "/" in export_id or "\\" in export_id:
            raise ExportFileMissingOrCorruptException(export_id, "Invalid characters in export_id.")

        managed_file_key = self.persistence.get_export_managed_key(forecast_run_id, export_id)
        file_path = self.exports_dir / managed_file_key

        try:
            resolved_path = file_path.resolve()
            resolved_dir = self.exports_dir.resolve()
            resolved_path.relative_to(resolved_dir)
        except (ValueError, RuntimeError):
            raise ExportFileMissingOrCorruptException(
                export_id, "Export path lies outside managed export storage."
            )

        if not resolved_path.exists() or not resolved_path.is_file():
            raise ExportFileMissingOrCorruptException(
                export_id, f"Export file missing on server disk at '{managed_file_key}'"
            )

        return resolved_path


# Module-level default export service
export_service = ExportService()
