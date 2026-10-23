import os
import sys
import time
from datetime import datetime
import openpyxl
import mysql.connector

# Cấu hình UTF-8 cho Windows console
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

DB_CONFIG = {
    'host': os.getenv('MYSQL_HOST', '127.0.0.1'),
    'port': int(os.getenv('MYSQL_PORT', 3306)),
    'user': os.getenv('MYSQL_USER', 'root'),
    'password': os.getenv('MYSQL_PASSWORD', '123456'),
    'database': os.getenv('MYSQL_DATABASE', 'online_retail_db'),
    'autocommit': False
}

def parse_row(row, year_cohort):
    """
    row: (Invoice, StockCode, Description, Quantity, InvoiceDate, Price, Customer ID, Country)
    """
    invoice, stock_code, description, quantity, invoice_date, price, customer_id, country = row

    if invoice is None or stock_code is None:
        return None

    # Invoice
    invoice_str = str(invoice).strip()
    if not invoice_str:
        return None

    # StockCode
    stock_code_str = str(stock_code).strip()

    # Description
    desc_str = str(description).strip()[:255] if description else None

    # Quantity
    try:
        qty = int(quantity) if quantity is not None else 0
    except (ValueError, TypeError):
        qty = 0

    # InvoiceDate
    if isinstance(invoice_date, datetime):
        dt = invoice_date
    elif invoice_date:
        try:
            dt = datetime.fromisoformat(str(invoice_date))
        except Exception:
            dt = None
    else:
        dt = None

    if dt is None:
        return None

    # Price
    try:
        unit_price = float(price) if price is not None else 0.0
    except (ValueError, TypeError):
        unit_price = 0.0

    # Customer ID
    cid = None
    if customer_id is not None:
        try:
            cid_float = float(customer_id)
            if not (cid_float != cid_float):  # check NaN
                cid = int(cid_float)
        except (ValueError, TypeError):
            cid = None

    # Country
    country_str = str(country).strip()[:50] if country else 'Unknown'

    return (invoice_str, stock_code_str, desc_str, qty, dt, unit_price, cid, country_str, year_cohort)


def import_excel_to_mysql(excel_path: str, batch_size: int = 10000):
    if not os.path.exists(excel_path):
        print(f"[!] Không tìm thấy file excel tại: {excel_path}")
        return

    print("=" * 60)
    print(" BẮT ĐẦU IMPORT DỮ LIỆU ONLINE RETAIL II VÀO MYSQL")
    print(f" File: {excel_path}")
    print(f" Database: {DB_CONFIG['database']} @ {DB_CONFIG['host']}:{DB_CONFIG['port']}")
    print("=" * 60)

    start_total_time = time.time()
    conn = mysql.connector.connect(**DB_CONFIG)
    cursor = conn.cursor()

    # Tối ưu tốc độ insert InnoDB
    cursor.execute("SET autocommit = 0;")
    cursor.execute("SET unique_checks = 0;")
    cursor.execute("SET foreign_key_checks = 0;")

    # Đảm bảo bảng tồn tại
    create_table_sql = """
    CREATE TABLE IF NOT EXISTS online_retail (
        id BIGINT AUTO_INCREMENT PRIMARY KEY,
        invoice_no VARCHAR(20) NOT NULL,
        stock_code VARCHAR(30) NOT NULL,
        description VARCHAR(255),
        quantity INT NOT NULL,
        invoice_date DATETIME NOT NULL,
        unit_price DECIMAL(10, 2) NOT NULL,
        customer_id INT NULL,
        country VARCHAR(50),
        year_cohort VARCHAR(20),
        INDEX idx_invoice_no (invoice_no),
        INDEX idx_customer_id (customer_id),
        INDEX idx_invoice_date (invoice_date),
        INDEX idx_stock_code (stock_code),
        INDEX idx_country (country)
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;
    """
    cursor.execute(create_table_sql)
    conn.commit()

    # Kiểm tra số dòng hiện tại
    cursor.execute("SELECT COUNT(*) FROM online_retail;")
    current_count = cursor.fetchone()[0]
    if current_count > 0:
        print(f"[*] Bảng online_retail hiện đã có {current_count:,} dòng.")
        # Nếu đã có dữ liệu, hỏi hoặc xóa dữ liệu cũ
        choice = os.getenv('TRUNCATE_TABLE', 'yes')
        if choice.lower() == 'yes':
            print("[-] Đang xóa dữ liệu cũ (TRUNCATE TABLE) để nạp mới toàn bộ...")
            cursor.execute("TRUNCATE TABLE online_retail;")
            conn.commit()
            print("[+] Đã xóa dữ liệu cũ thành công.")

    insert_sql = """
    INSERT INTO online_retail (
        invoice_no, stock_code, description, quantity, invoice_date, unit_price, customer_id, country, year_cohort
    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
    """

    print("[*] Đang mở file Excel (read-only stream)...")
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    total_imported = 0

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        year_cohort = '2009-2010' if '2009' in sheet_name else '2010-2011'
        print(f"\n>>> Đang xử lý Sheet: '{sheet_name}' (Cohort: {year_cohort}) ...")

        batch = []
        sheet_count = 0
        sheet_start_time = time.time()

        for idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if idx == 0:
                # Dòng tiêu đề header
                continue

            parsed = parse_row(row, year_cohort)
            if parsed:
                batch.append(parsed)

            if len(batch) >= batch_size:
                cursor.executemany(insert_sql, batch)
                conn.commit()
                sheet_count += len(batch)
                total_imported += len(batch)
                batch.clear()
                elapsed = time.time() - sheet_start_time
                rate = sheet_count / elapsed if elapsed > 0 else 0
                print(f"    -> Đã nạp: {sheet_count:,} dòng (Tốc độ: {rate:,.0f} dòng/giây)...")

        # Nạp nốt các dòng còn lại của sheet
        if batch:
            cursor.executemany(insert_sql, batch)
            conn.commit()
            sheet_count += len(batch)
            total_imported += len(batch)
            batch.clear()

        sheet_elapsed = time.time() - sheet_start_time
        print(f"[+] Hoàn thành Sheet '{sheet_name}': {sheet_count:,} dòng trong {sheet_elapsed:.1f}s.")

    wb.close()

    # Bật lại các kiểm tra
    cursor.execute("SET unique_checks = 1;")
    cursor.execute("SET foreign_key_checks = 1;")
    conn.commit()

    total_elapsed = time.time() - start_total_time
    print("\n" + "=" * 60)
    print(f" IMPORT HOÀN TẤT THÀNH CÔNG!")
    print(f" Tổng số dòng đã nạp vào MySQL: {total_imported:,}")
    print(f" Tổng thời gian thực hiện: {total_elapsed:.1f} giây")
    print("=" * 60)

    # Hiển thị thống kê nhanh từ CSDL
    cursor.execute("""
    SELECT 
        COUNT(*) as total_rows,
        COUNT(DISTINCT invoice_no) as total_invoices,
        COUNT(DISTINCT customer_id) as total_customers,
        COUNT(DISTINCT stock_code) as total_products,
        MIN(invoice_date) as start_date,
        MAX(invoice_date) as end_date
    FROM online_retail;
    """)
    stats = cursor.fetchone()
    print("\n THỐNG KÊ NHANH DỮ LIỆU ĐÃ NẠP:")
    print(f"  - Tổng số bản ghi (rows): {stats[0]:,}")
    print(f"  - Tổng số hóa đơn (invoices): {stats[1]:,}")
    print(f"  - Tổng số khách hàng (customers): {stats[2]:,}")
    print(f"  - Tổng số mặt hàng (products): {stats[3]:,}")
    print(f"  - Khoảng thời gian: {stats[4]}  -->  {stats[5]}")

    cursor.close()
    conn.close()

if __name__ == '__main__':
    default_excel = os.path.join(os.path.dirname(__file__), '..', '..', 'online_retail_II.xlsx')
    if len(sys.argv) > 1:
        default_excel = sys.argv[1]
    
    excel_abs = os.path.abspath(default_excel)
    import_excel_to_mysql(excel_abs, batch_size=10000)
